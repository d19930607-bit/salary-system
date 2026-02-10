import streamlit as st
import pandas as pd
from datetime import datetime
import gspread
from oauth2client.service_account import ServiceAccountCredentials
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import io
import time

# --- 設定檔 ---
JSON_FILE = "service_account.json"
SHEET_NAME = "工務薪資系統_資料庫"
ALLOWANCE_PER_DAY = 200

# --- 連線設定 ---
@st.cache_resource
def init_connection():
    scope = ["https://spreadsheets.google.com/feeds", "https://www.googleapis.com/auth/drive"]
    creds = ServiceAccountCredentials.from_json_keyfile_name(JSON_FILE, scope)
    client = gspread.authorize(creds)
    return client

def get_sheet(client, name):
    try:
        return client.open(SHEET_NAME).worksheet(name)
    except:
        st.cache_resource.clear()
        return init_connection().open(SHEET_NAME).worksheet(name)

# --- 資料讀取 ---
def load_data():
    client = init_connection()
    s_emp = get_sheet(client, "員工資料")
    s_work = get_sheet(client, "工作紀錄")
    
    emp_rows = s_emp.get_all_values()
    employees = [r[0] for r in emp_rows[1:] if r]
    salary_map = {r[0]: int(r[1]) for r in emp_rows[1:] if len(r)>1 and r[1].isdigit()}
    
    work_rows = s_work.get_all_values()
    
    locations = sorted(list(set([r[2] for r in work_rows[1:] if len(r)>2 and r[2]])))
    work_details = sorted(list(set([r[6] for r in work_rows[1:] if len(r)>8 and r[8]=="手動登錄" and r[6]])))
    money_items = sorted(list(set([r[6] for r in work_rows[1:] if len(r)>8 and r[8] in ["扣款", "獎金"] and r[6]])))
    
    return employees, salary_map, locations, work_details, money_items, work_rows, s_work, s_emp, emp_rows

# --- Excel 報表邏輯 ---
def generate_excel(year, work_rows, employees):
    data = [r[:9] for r in work_rows[1:] if len(r)>=9]
    df = pd.DataFrame(data, columns=['日期','姓名','地點','工時','日薪','津貼','備註','金額','類型'])
    df['日期'] = pd.to_datetime(df['日期'], errors='coerce')
    for c in ['工時','日薪','津貼','金額']: df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0)
    df.loc[df['類型'] == '月結輸入', '類型'] = '扣款'

    wb = Workbook()
    ft_head = Font(name='微軟正黑體', size=14, bold=True)
    ft_bold = Font(name='微軟正黑體', bold=True)
    ft_red = Font(name='微軟正黑體', bold=True, color="FF0000")
    ft_blue = Font(name='微軟正黑體', bold=True, color="0000FF")
    bd = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    fill_head = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    align_c = Alignment(horizontal="center", vertical="center")

    # 1. 年度總表
    ws_y = wb.active; ws_y.title = "年度總表"
    ws_y.merge_cells("A1:F1"); c=ws_y.cell(1,1,"【年度員工薪資總結】"); c.font=ft_head
    headers = ["姓名", "總天數", "總工資", "總津貼", "總扣款", "實支總額"]
    for i, h in enumerate(headers, 1):
        c = ws_y.cell(2, i, h); c.font = ft_bold; c.border = bd; c.fill = fill_head; c.alignment = align_c
        ws_y.column_dimensions[get_column_letter(i)].width = 15
    
    df_year = df[df['日期'].dt.year == int(year)]
    yr_row = 3
    for p in employees:
        if p not in df_year['姓名'].values: continue
        p_d = df_year[df_year['姓名'] == p]
        work = p_d[p_d['類型'] == '手動登錄']; deduct = p_d[p_d['類型'] == '扣款']; bonus = p_d[p_d['類型'] == '獎金']
        t_d = work['工時'].sum(); w_s = (work['工時']*work['日薪']).sum(); f_s = (work['工時']*ALLOWANCE_PER_DAY).sum()
        d_s = deduct['金額'].sum(); b_s = bonus['金額'].sum(); net = w_s + f_s + b_s - d_s
        vals = [p, t_d, w_s, f_s + b_s, d_s, net]
        for i, v in enumerate(vals, 1): ws_y.cell(yr_row, i, v).border = bd; ws_y.cell(yr_row, i).alignment = align_c
        yr_row += 1
    
    site_row = yr_row + 2
    ws_y.merge_cells(f"A{site_row}:D{site_row}"); c=ws_y.cell(site_row,1,"【年度案場總結】"); c.font=ft_head
    site_heads = ["案場名稱", "總施工天數", "總出工人數", "總人次(工)"]
    for i, h in enumerate(site_heads, 1):
        c = ws_y.cell(site_row+1, i, h); c.font = ft_bold; c.border = bd; c.fill = fill_head; c.alignment = align_c
        ws_y.column_dimensions[get_column_letter(i)].width = 18
    
    site_stats = df_year[df_year['類型']=='手動登錄'].groupby('地點').agg({'日期':'nunique', '姓名':'nunique', '工時':'sum'}).reset_index()
    site_data_row = site_row + 2
    for _, r in site_stats.iterrows():
        vals = [r['地點'], r['日期'], r['姓名'], r['工時']]
        for i, v in enumerate(vals, 1): ws_y.cell(site_data_row, i, v).border = bd; ws_y.cell(site_data_row, i).alignment = align_c
        site_data_row += 1

    # 2. 月份報表
    for m in range(1, 13):
        df_m = df[(df['日期'].dt.year == int(year)) & (df['日期'].dt.month == m)]
        if df_m.empty: continue
        ws = wb.create_sheet(f"{m}月")
        active_emp = [p for p in employees if p in df_m['姓名'].unique()]
        
        # A. 薪資明細
        ws.merge_cells("A1:G1"); c=ws.cell(1,1,"【薪資明細】"); c.font=ft_head
        month_bonuses = sorted(df_m[df_m['類型']=='獎金']['備註'].unique())
        month_deducts = sorted(df_m[df_m['類型']=='扣款']['備註'].unique())
        month_adjusts = sorted(df_m[df_m['類型']=='薪資調整']['備註'].unique())
        all_labels = ["總天數", "日薪", "工資小計", "便當飲料", "小計"] + month_bonuses + month_deducts + month_adjusts + ["實支總額"]
        
        for i, lab in enumerate(all_labels):
            r = i + 3; c = ws.cell(r, 1, lab); c.border=bd; c.alignment=align_c
            if lab in ["便當飲料", "小計"] or lab in month_bonuses: c.font = ft_blue
            if lab in month_deducts: c.font = ft_red
            if lab == "實支總額": c.font = ft_bold
            if lab in month_adjusts: c.font = ft_bold
        ws.column_dimensions['A'].width = 20
        
        curr_col = 2
        col_map = {}
        for p in active_emp:
            ws.merge_cells(start_row=2, start_column=curr_col, end_row=2, end_column=curr_col+1)
            c_head = ws.cell(2, curr_col, f"{int(year)-1911}年{m}月   {p}")
            c_head.font = ft_bold; c_head.alignment = align_c; c_head.border = bd; ws.cell(2, curr_col+1).border = bd
            
            p_data = df_m[df_m['姓名'] == p]
            work = p_data[p_data['類型'] == '手動登錄']
            deduct_df = p_data[p_data['類型'] == '扣款']
            bonus_df = p_data[p_data['類型'] == '獎金']
            adjust_df = p_data[p_data['類型'] == '薪資調整']
            
            t_days = work['工時'].sum(); t_rate = work['日薪'].max()
            w_sum = (work['工時']*work['日薪']).sum(); f_s = (work['工時']*ALLOWANCE_PER_DAY).sum()
            sub_1 = w_sum + f_s
            
            vals = [t_days, t_rate, w_sum, f_s, sub_1]
            styles = [None, None, None, ft_blue, ft_blue]
            current_r = 3
            for v, s in zip(vals, styles):
                ws.merge_cells(start_row=current_r, start_column=curr_col, end_row=current_r, end_column=curr_col+1)
                c_val = ws.cell(current_r, curr_col, v); c_val.border = bd; c_val.alignment = align_c
                if s: c_val.font = s
                ws.cell(current_r, curr_col+1).border = bd
                current_r += 1
                
            for b_item in month_bonuses:
                amt = bonus_df[bonus_df['備註']==b_item]['金額'].sum()
                val = amt if amt > 0 else ""
                ws.merge_cells(start_row=current_r, start_column=curr_col, end_row=current_r, end_column=curr_col+1)
                c_val = ws.cell(current_r, curr_col, val); c_val.border = bd; c_val.alignment = align_c; c_val.font = ft_blue
                ws.cell(current_r, curr_col+1).border = bd; current_r += 1
            
            for d_item in month_deducts:
                amt = deduct_df[deduct_df['備註']==d_item]['金額'].sum()
                val = amt if amt > 0 else ""
                ws.merge_cells(start_row=current_r, start_column=curr_col, end_row=current_r, end_column=curr_col+1)
                c_val = ws.cell(current_r, curr_col, val); c_val.border = bd; c_val.alignment = align_c; c_val.font = ft_red
                ws.cell(current_r, curr_col+1).border = bd; current_r += 1

            for a_item in month_adjusts:
                has_adj = not adjust_df[adjust_df['備註']==a_item].empty
                val = "✔️" if has_adj else ""
                ws.merge_cells(start_row=current_r, start_column=curr_col, end_row=current_r, end_column=curr_col+1)
                c_val = ws.cell(current_r, curr_col, val); c_val.border = bd; c_val.alignment = align_c
                ws.cell(current_r, curr_col+1).border = bd; current_r += 1
            
            net = sub_1 + bonus_df['金額'].sum() - deduct_df['金額'].sum()
            ws.merge_cells(start_row=current_r, start_column=curr_col, end_row=current_r, end_column=curr_col+1)
            c_net = ws.cell(current_r, curr_col, net); c_net.border = bd; c_net.alignment = align_c; c_net.font = ft_bold
            ws.cell(current_r, curr_col+1).border = bd
            col_map[p] = curr_col; curr_col += 2
            
        # B. 出勤明細
        start_row = 3 + len(all_labels) + 2
        ws.merge_cells(f"A{start_row}:C{start_row}"); c=ws.cell(start_row,1,"【出勤明細】"); c.font=ft_head
        work_only = df_m[df_m['類型']=='手動登錄']
        dates = sorted(work_only['日期'].unique())
        
        ws.cell(start_row+1, 1, "日期").border=bd
        for p in active_emp:
            ws.merge_cells(start_row=start_row+1, start_column=col_map[p], end_row=start_row+1, end_column=col_map[p]+1)
            c = ws.cell(start_row+1, col_map[p], p); c.border=bd; c.alignment=align_c
            ws.cell(start_row+1, col_map[p]+1).border=bd
        
        curr_r = start_row + 2
        for d in dates:
            ws.cell(curr_r, 1, d.strftime("%m月%d日")).border=bd
            day_data = work_only[work_only['日期'] == d]
            for p in active_emp:
                rec = day_data[day_data['姓名'] == p]
                txt = ""
                if not rec.empty:
                    locs = rec['地點'].tolist(); dys = rec['工時'].tolist()
                    txt = "\n".join([f"{l} ({d})" for l, d in zip(locs, dys)])
                ws.merge_cells(start_row=curr_r, start_column=col_map[p], end_row=curr_r, end_column=col_map[p]+1)
                c = ws.cell(curr_r, col_map[p], txt); c.border=bd; c.alignment=Alignment(wrap_text=True, horizontal='center', vertical='center')
                ws.cell(curr_r, col_map[p]+1).border=bd
            curr_r += 1
            
        # C. 案場統計 & D. 工項統計
        i_col = 9
        ws.merge_cells(start_row=1, start_column=i_col, end_row=1, end_column=i_col+3); c=ws.cell(1, i_col, "【案場統計】"); c.font=ft_head
        for idx, h in enumerate(["案場", "天數", "人數", "工數"]): ws.cell(2, i_col+idx, h).border=bd
        site_stats = work_only.groupby('地點').agg({'日期':'nunique', '姓名':'nunique', '工時':'sum'}).reset_index()
        for idx, row in site_stats.iterrows():
            for c_idx, val in enumerate([row['地點'], row['日期'], row['姓名'], row['工時']]): ws.cell(idx+3, i_col+c_idx, val).border=bd

        n_col = 14
        ws.merge_cells(start_row=1, start_column=n_col, end_row=1, end_column=n_col+4); c=ws.cell(1, n_col, "【案場工項統計】"); c.font=ft_head
        for idx, h in enumerate(["日期", "案場", "工項", "工數", "人數"]): ws.cell(2, n_col+idx, h).border=bd
        work_only['備註'] = work_only['備註'].replace("", "一般")
        det_stats = work_only.groupby(['日期', '地點', '備註']).agg({'工時':'sum', '姓名':'nunique'}).reset_index().sort_values('日期')
        for idx, row in det_stats.iterrows():
            vals = [row['日期'].strftime("%m/%d"), row['地點'], row['備註'], row['工時'], row['姓名']]
            for c_idx, val in enumerate(vals): ws.cell(idx+3, n_col+c_idx, val).border=bd

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# --- 主介面 ---
def main():
    st.set_page_config(page_title="工務薪資系統 (Web)", page_icon="🏗️")
    st.title("🏗️ 工務薪資系統")

    menu = st.selectbox("≡ 功能選單", ["📅 工務登錄", "💰 獎金與扣款", "🛠️ 員工管理", "💵 薪資明細瀏覽", "🏗️ 案場統計瀏覽", "📊 報表下載"])
    
    employees, salary_map, locations, work_details, money_items, work_rows, s_work, s_emp, emp_rows = load_data()
    client = init_connection()

    if menu == "📅 工務登錄":
        st.subheader("每日出勤登錄")
        
        with st.form("work_form", clear_on_submit=True):
            st.info("💡 下拉選單找不到時，請直接在下方欄位打字。")
            date = st.date_input("📅 日期", datetime.now())
            emp = st.selectbox("👷 員工", employees)
            
            loc_select = st.selectbox("📍 地點 (選擇)", [""] + locations)
            loc_input = st.text_input("📍 地點 (手動輸入)")
            
            detail_select = st.selectbox("🔧 工項 (選擇)", [""] + work_details)
            detail_input = st.text_input("🔧 工項 (手動輸入)")

            hours = st.radio("⏰ 工時", [1.0, 0.5], horizontal=True)
            
            if st.form_submit_button("💾 儲存紀錄", type="primary"):
                final_loc = loc_input if loc_input else loc_select
                final_detail = detail_input if detail_input else detail_select
                
                if not final_loc:
                    st.error("❌ 請選擇或輸入地點！")
                else:
                    wage = salary_map.get(emp, 0)
                    allowance = int(hours * ALLOWANCE_PER_DAY)
                    row = [str(date), emp, final_loc, hours, wage, allowance, final_detail, 0, "手動登錄"]
                    s_work.append_row(row)
                    st.success(f"✅ 已儲存：{emp} @ {final_loc}")
                    st.cache_resource.clear()
                    time.sleep(1)
                    st.rerun()

        st.divider()
        st.subheader("📋 最近 10 筆紀錄 (最新在上)") # [V73 修正] 強制顯示列表
        
        # 顯示最近 10 筆 (降冪)
        df_work = pd.DataFrame(work_rows[1:], columns=work_rows[0])
        df_work = df_work[df_work['類型'] == '手動登錄']
        df_work['日期'] = pd.to_datetime(df_work['日期'], errors='coerce')
        df_work = df_work.sort_values(by='日期', ascending=False)
        
        st.dataframe(
            df_work[['日期', '姓名', '地點', '備註', '工時(天)']].head(10).style.format({"日期": lambda t: t.strftime("%Y-%m-%d")}),
            hide_index=True
        )
        
        with st.expander("🛠️ 修改或刪除舊紀錄"):
            options = []
            for i in range(len(work_rows)-1, 0, -1):
                r = work_rows[i]
                if len(r) > 8 and r[8] == "手動登錄":
                    display = f"{r[0]} | {r[1]} | {r[2]}"
                    options.append((i + 1, display))
            
            if not options:
                st.warning("沒有可修改的紀錄")
            else:
                selected_opt = st.selectbox("選擇要處理的紀錄", options, format_func=lambda x: x[1])
                target_row_idx = selected_opt[0]
                target_row_data = work_rows[target_row_idx - 1]
                
                with st.form("edit_work_form"):
                    e_date = st.date_input("日期", datetime.strptime(target_row_data[0], "%Y-%m-%d"))
                    e_emp = st.selectbox("員工", employees, index=employees.index(target_row_data[1]) if target_row_data[1] in employees else 0)
                    e_loc = st.text_input("地點", value=target_row_data[2])
                    e_detail = st.text_input("工項", value=target_row_data[6])
                    e_hours = st.radio("工時", [1.0, 0.5], index=0 if float(target_row_data[3])==1.0 else 1)
                    
                    c1, c2 = st.columns(2)
                    if c1.form_submit_button("✅ 確認修改"):
                        wage = salary_map.get(e_emp, 0)
                        allowance = int(e_hours * ALLOWANCE_PER_DAY)
                        new_row = [str(e_date), e_emp, e_loc, e_hours, wage, allowance, e_detail, 0, "手動登錄"]
                        s_work.update(f"A{target_row_idx}:I{target_row_idx}", [new_row])
                        st.success("修改完成！")
                        st.cache_resource.clear()
                        time.sleep(1)
                        st.rerun()
                        
                    if c2.form_submit_button("🗑️ 刪除紀錄", type="primary"):
                        s_work.delete_rows(target_row_idx)
                        st.warning("已刪除紀錄")
                        st.cache_resource.clear()
                        time.sleep(1)
                        st.rerun()

    elif menu == "💰 獎金與扣款":
        st.subheader("獎金與扣款管理")
        
        with st.form("money_form", clear_on_submit=True):
            m_type = st.radio("類型", ["扣款", "獎金"], horizontal=True)
            date = st.date_input("📅 日期", datetime.now())
            emp = st.selectbox("👷 對象", employees)
            item_select = st.selectbox("📝 項目 (選擇)", [""] + money_items)
            item_input = st.text_input("📝 項目 (手動輸入)")
            amt = st.number_input("💲 金額", min_value=0, step=100)
            
            if st.form_submit_button("💰 確認新增", type="primary"):
                final_item = item_input if item_input else item_select
                if not final_item:
                    st.error("❌ 請輸入項目名稱！")
                else:
                    row = [str(date), emp, "", 0, 0, 0, final_item, amt, m_type]
                    s_work.append_row(row)
                    st.success(f"✅ 已新增：{emp} {m_type} ${amt}")
                    st.cache_resource.clear()
                    time.sleep(1)
                    st.rerun()

        st.divider()
        with st.expander("🛠️ 修改或刪除 款項紀錄"):
            m_options = []
            for i in range(len(work_rows)-1, 0, -1):
                r = work_rows[i]
                if len(r) > 8 and r[8] in ["扣款", "獎金"]:
                    display = f"{r[0]} | {r[1]} | {r[8]} {r[6]} ${r[7]}"
                    m_options.append((i + 1, display))
            
            if not m_options:
                st.warning("沒有可修改的紀錄")
            else:
                sel_m = st.selectbox("選擇紀錄", m_options, format_func=lambda x: x[1])
                m_row_idx = sel_m[0]
                m_data = work_rows[m_row_idx - 1]
                
                with st.form("edit_money_form"):
                    st.write(f"編輯：{m_data[1]} 的 {m_data[6]}")
                    em_type = st.radio("類型", ["扣款", "獎金"], index=0 if m_data[8]=="扣款" else 1)
                    em_date = st.date_input("日期", datetime.strptime(m_data[0], "%Y-%m-%d"))
                    em_emp = st.selectbox("對象", employees, index=employees.index(m_data[1]) if m_data[1] in employees else 0)
                    em_item = st.text_input("項目", value=m_data[6])
                    em_amt = st.number_input("金額", value=int(m_data[7]), step=100)
                    
                    c1, c2 = st.columns(2)
                    if c1.form_submit_button("✅ 修改"):
                        new_row = [str(em_date), em_emp, "", 0, 0, 0, em_item, em_amt, em_type]
                        s_work.update(f"A{m_row_idx}:I{m_row_idx}", [new_row])
                        st.success("已更新")
                        st.cache_resource.clear()
                        time.sleep(1)
                        st.rerun()
                    if c2.form_submit_button("🗑️ 刪除", type="primary"):
                        s_work.delete_rows(m_row_idx)
                        st.warning("已刪除")
                        st.cache_resource.clear()
                        time.sleep(1)
                        st.rerun()

    elif menu == "🛠️ 員工管理":
        st.subheader("員工與薪資調整")
        
        with st.expander("➕ 新增員工"):
            with st.form("add_emp"):
                new_n = st.text_input("姓名")
                new_w = st.number_input("日薪", min_value=0, step=100)
                if st.form_submit_button("新增"):
                    s_emp.append_row([new_n, new_w])
                    st.success(f"✅ 員工 {new_n} 已新增")
                    st.cache_resource.clear()
                    st.rerun()

        st.divider()
        st.write("💲 **調整薪資 (含歷程)**")
        
        with st.form("adj_wage"):
            edit_emp = st.selectbox("選擇員工", employees)
            emp_row_idx = -1
            for i, r in enumerate(emp_rows):
                if len(r)>0 and r[0] == edit_emp:
                    emp_row_idx = i + 1
                    break
            
            new_wage = st.number_input("新日薪", min_value=0, step=100)
            eff_date = st.date_input("生效日期", datetime.now())
            
            if st.form_submit_button("更新並記錄"):
                current_wage = salary_map.get(edit_emp, 0)
                s_emp.update_cell(emp_row_idx, 2, new_wage)
                diff = new_wage - current_wage
                note = f"{eff_date}起 {'加薪' if diff>0 else '減薪'}{abs(diff)}"
                row = [str(eff_date), edit_emp, "", 0, 0, 0, note, 0, "薪資調整"]
                s_work.append_row(row)
                st.success(f"✅ 已更新 {edit_emp} 薪資！({note})")
                st.cache_resource.clear()
                time.sleep(1)
                st.rerun()
        
        st.divider()
        with st.expander("🗑️ 刪除員工"):
            del_emp_name = st.selectbox("選擇要刪除的員工", employees, key="del_sel")
            if st.button("確認刪除員工", type="primary"):
                for i, r in enumerate(emp_rows):
                    if len(r)>0 and r[0] == del_emp_name:
                        s_emp.delete_rows(i+1)
                        st.warning(f"已刪除 {del_emp_name}")
                        st.cache_resource.clear()
                        time.sleep(1)
                        st.rerun()
                        break

    # [V73 新功能] 薪資明細瀏覽
    elif menu == "💵 薪資明細瀏覽":
        st.subheader("💵 員工薪資統計")
        
        # 篩選區
        col1, col2 = st.columns(2)
        year = col1.selectbox("年份", [2025, 2026, 2027], index=1)
        mode = col2.radio("統計範圍", ["本月", "全年度"], horizontal=True)
        
        if mode == "本月":
            month = st.selectbox("月份", range(1, 13), index=datetime.now().month-1)
        
        # 處理資料
        data = [r[:9] for r in work_rows[1:] if len(r)>=9]
        df = pd.DataFrame(data, columns=['日期','姓名','地點','工時','日薪','津貼','備註','金額','類型'])
        df['日期'] = pd.to_datetime(df['日期'], errors='coerce')
        for c in ['工時','日薪','津貼','金額']: df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0)
        df.loc[df['類型'] == '月結輸入', '類型'] = '扣款'
        
        # 篩選日期
        if mode == "本月":
            df = df[(df['日期'].dt.year == int(year)) & (df['日期'].dt.month == int(month))]
        else:
            df = df[df['日期'].dt.year == int(year)]
            
        if df.empty:
            st.warning("查無資料")
        else:
            # 計算每人薪資
            salary_data = []
            active_emps = df['姓名'].unique()
            for p in active_emps:
                p_df = df[df['姓名'] == p]
                work = p_df[p_df['類型'] == '手動登錄']
                deduct = p_df[p_df['類型'] == '扣款']['金額'].sum()
                bonus = p_df[p_df['類型'] == '獎金']['金額'].sum()
                
                days = work['工時'].sum()
                base = (work['工時'] * work['日薪']).sum()
                allow = (work['工時'] * ALLOWANCE_PER_DAY).sum()
                total = base + allow + bonus - deduct
                
                salary_data.append({
                    "姓名": p, "出勤天數": days, "應支(含津貼/獎金)": base+allow+bonus, "扣款": deduct, "實領總額": total
                })
            
            st.dataframe(pd.DataFrame(salary_data), hide_index=True)

    # [V73 新功能] 案場統計瀏覽
    elif menu == "🏗️ 案場統計瀏覽":
        st.subheader("🏗️ 案場與工項統計")
        
        col1, col2 = st.columns(2)
        year = col1.selectbox("年份", [2025, 2026, 2027], index=1, key="site_y")
        mode = col2.radio("統計範圍", ["本月", "全年度"], horizontal=True, key="site_m")
        
        if mode == "本月":
            month = st.selectbox("月份", range(1, 13), index=datetime.now().month-1, key="site_mon")
            
        data = [r[:9] for r in work_rows[1:] if len(r)>=9]
        df = pd.DataFrame(data, columns=['日期','姓名','地點','工時','日薪','津貼','備註','金額','類型'])
        df['日期'] = pd.to_datetime(df['日期'], errors='coerce')
        df = df[df['類型'] == '手動登錄']
        for c in ['工時']: df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0)
        
        if mode == "本月":
            df = df[(df['日期'].dt.year == int(year)) & (df['日期'].dt.month == int(month))]
        else:
            df = df[df['日期'].dt.year == int(year)]
            
        if df.empty:
            st.warning("查無資料")
        else:
            st.write("📊 **案場總表** (哪個工地做最多)")
            site_stats = df.groupby('地點').agg({'工時':'sum', '姓名':'nunique', '日期':'nunique'}).reset_index()
            site_stats.columns = ['案場', '總工數', '出工人數', '施工天數']
            st.dataframe(site_stats, hide_index=True)
            
            st.divider()
            st.write("🔧 **工項細目** (都在做什麼)")
            df['備註'] = df['備註'].replace("", "一般出勤")
            detail_stats = df.groupby(['地點', '備註']).agg({'工時':'sum'}).reset_index()
            detail_stats.columns = ['案場', '工項', '總工數']
            st.dataframe(detail_stats, hide_index=True)

    elif menu == "📊 報表下載":
        st.subheader("Excel 報表匯出")
        year = st.selectbox("選擇年份", [2025, 2026, 2027], index=1)
        
        if st.button("🚀 產生報表", type="primary"):
            with st.spinner("正在生成 Excel..."):
                excel_data = generate_excel(year, work_rows, employees)
                st.success("生成完畢！")
                st.download_button(
                    label="📥 點此下載檔案",
                    data=excel_data,
                    file_name=f"工務薪資報表_{year}年.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

if __name__ == "__main__":
    main()
